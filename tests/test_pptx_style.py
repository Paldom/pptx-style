#!/usr/bin/env python3
"""Acceptance tests for pptx-style-profile + pptx-style-build (stdlib; no LibreOffice needed).

    python3 tests/test_pptx_style.py

Generates the brand-free fixtures, then checks the acceptance criteria: profiling three
different templates (one messy), the drift diff, a 18-slide round trip with cloned
patterns and notes, off-contract rejection, overflow and split, edit-mode passthrough,
per-clone chart data, package hygiene, and a clean skip when LibreOffice is absent.
python-pptx is used only if importable (dev tooling), to prove the output opens there.
"""

from __future__ import annotations

import contextlib
import importlib.util
import io
import json
import os
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PROFILE = ROOT / "skills" / "pptx-style-profile" / "scripts"
BUILD = ROOT / "skills" / "pptx-style-build" / "scripts"
os.environ["PPTX_STYLE_FONT_DIRS"] = str(ROOT / "tests" / "fonts")
sys.path.insert(0, str(BUILD))

import build  # noqa: E402
import lint  # noqa: E402
import pptx_core as core  # noqa: E402
import render  # noqa: E402


def load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


profile_cli = load(PROFILE / "profile.py", "profile_cli")
diff_cli = load(PROFILE / "diff.py", "diff_cli")
fixtures = load(ROOT / "tests" / "fixtures" / "make_fixtures.py", "make_fixtures")

TMP = Path(tempfile.mkdtemp(prefix="pptx-style-test-"))
fixtures.main(["x", str(TMP)])


def quiet(fn, *a, **kw):
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = fn(*a, **kw)
    return rc, out.getvalue() + err.getvalue()


def contract_for(name: str) -> Path:
    out = TMP / ".pptx-style" / name
    if not (out / "contract.json").exists():
        rc, log = quiet(profile_cli.main, [str(TMP / name), "--out", str(out)])
        assert rc == 0, log
    return out / "contract.json"


def run_spec(spec: dict, name: str, *extra):
    p = TMP / f"{name}.json"
    spec.setdefault("output", f"{name}.pptx")
    p.write_text(json.dumps(spec))
    return quiet(build.main, [str(p), *extra])


def agenda_fill():
    f = {"title": "Agenda"}
    for i in range(5):
        f[["body", "body3", "body5", "body7", "body9"][i]] = f"0{i + 1}"
        f[["body2", "body4", "body6", "body8", "body10"][i]] = f"Topic {chr(65 + i)} in brief"
    return f


def round_trip_spec(contract: Path) -> dict:
    s = []

    def add(slide, notes=True):
        if notes:
            slide["notes"] = f"Speaker notes for slide {len(s) + 1}: say why it matters."
        s.append(slide)

    add(
        {
            "layout": "Címdia",
            "fill": {"title": "Field operations review", "subtitle": "Quarterly update"},
        },
        notes=False,
    )
    add({"layout": "Agenda", "fill": agenda_fill()})
    add(
        {
            "layout": "Section Header",
            "fill": {"title": "Where we are", "body": "Context and numbers"},
        },
        notes=False,
    )
    add(
        {
            "layout": "Title and Content",
            "fill": {
                "title": "Demand doubled in two regions",
                "body": ["North grew fastest", "South steady", "Capacity is the constraint"],
            },
        }
    )
    add(
        {
            "pattern": "key-numbers",
            "fill": {
                "title": "Three numbers that matter",
                "t1": "Source: internal tracker",
                "g1": [
                    {"t1": "2x", "t2": "orders"},
                    {"t1": "14 d", "t2": "lead time"},
                    {"t1": "93%", "t2": "on-time rate"},
                ],
            },
        }
    )
    add(
        {
            "pattern": "three-ways-we-help",
            "fill": {
                "t1": "Three levers",
                "g1": [
                    {"t1": "Hire", "t2": "Add a second shift."},
                    {"t1": "Automate", "t2": "Pick-to-light in hall B."},
                ],
            },
        }
    )
    add(
        {
            "layout": "Two Content",
            "fill": {
                "title": "Options side by side",
                "body": ["Build in-house", "Slower start"],
                "body2": ["Partner", "Faster, less control"],
            },
        }
    )
    add(
        {
            "pattern": "option-comparison",
            "fill": {
                "title": "Partner wins on speed",
                "tbl1": [
                    ["Criterion", "In-house", "Partner"],
                    ["Speed", "", ""],
                    ["Cost", "", ""],
                    ["Risk", "", ""],
                ],
                "mark18": None,
            },
        }
    )
    add(
        {"layout": "Section Header", "fill": {"title": "The plan", "body": "Four quarters"}},
        notes=False,
    )
    add(
        {
            "pattern": "timeline",
            "fill": {
                "title": "Rollout over four quarters",
                "tbl1": [
                    ["Stream", "Q1", "Q2", "Q3", "Q4"],
                    ["Pilot", "", "", "", ""],
                    ["Scale", "", "", "", ""],
                    ["Run", "", "", "", ""],
                ],
                "bar1": {"row": 2, "from": 2, "to": 2},
                "bar2": {"row": 3, "from": 3, "to": 4},
                "bar3": {"row": 4, "from": 5, "to": 5},
            },
        }
    )
    add(
        {
            "layout": "Title and Content",
            "fill": {
                "title": "Risks we are watching",
                "body": [
                    "Hiring lead time",
                    "Vendor lock-in",
                    {"text": "Exit clause negotiated", "level": 1},
                ],
            },
        }
    )
    add({"layout": "Title Only", "fill": {"title": "Budget is unchanged"}})
    add(
        {
            "pattern": "key-numbers",
            "fill": {
                "title": "What success looks like",
                "t1": "Targets for Q4",
                "g1": [{"t1": "<10 d", "t2": "lead time"}, {"t1": "97%", "t2": "on time"}],
            },
        }
    )
    add(
        {
            "layout": "Title and Content",
            "fill": {
                "title": "Decisions needed today",
                "body": ["Approve the partner pilot", "Release hiring budget"],
            },
        }
    )
    add(
        {
            "pattern": "three-ways-we-help",
            "fill": {
                "t1": "Next steps",
                "g1": [
                    {"t1": "Week 1", "t2": "Sign the pilot."},
                    {"t1": "Week 2", "t2": "Start hiring."},
                    {"t1": "Week 4", "t2": "First review."},
                ],
            },
        }
    )
    add(
        {
            "layout": "Title and Content",
            "fill": {
                "title": "Appendix: method",
                "body": ["Data from the order tracker", "Q1-Q3 actuals"],
            },
        }
    )
    add({"layout": "Title Only", "fill": {"title": "Questions"}})
    add(
        {"layout": "Címdia", "fill": {"title": "Thank you", "subtitle": "ops-team@example.org"}},
        notes=False,
    )
    return {"contract": str(contract), "slides": s}


# --- tests ----------------------------------------------------------------------------


def test_core_copies_identical():
    assert (PROFILE / "pptx_core.py").read_bytes() == (BUILD / "pptx_core.py").read_bytes(), (
        "sync pptx_core.py"
    )


def test_profile_three_templates():
    for name in ("clean.pptx", "messy.potx", "charts.pptx"):
        c = core.load_contract(contract_for(name))
        assert c["layouts"] and c["patterns"] and c["theme"]["colors"], name
        assert (contract_for(name).parent / "STYLE.md").read_text().startswith("# Style guide")
        assert len(c["source"]["sha256"]) == 64 and c["source"]["parts"]
    c = core.load_contract(contract_for("clean.pptx"))
    assert {p["id"] for p in c["patterns"]} >= {
        "three-ways-we-help",
        "timeline",
        "option-comparison",
        "key-numbers",
    }
    assert any(ch["role"] == "slide_number" for m in c["masters"] for ch in m["chrome"])


def test_profile_messy_traits():
    c = core.load_contract(contract_for("messy.potx"))
    codes = {a["code"] for a in c["ai_readiness"]}
    for need in (
        "prompt-text-as-content",
        "duplicate-shape-names",
        "export-artifact-names",
        "hardcoded-colour",
        "localized-layout-name",
        "decorative-sample-slide",
    ):
        assert need in codes, need
    assert c["source"]["kind"] == "potx"
    title = next(L for L in c["layouts"] if L["name"] == "Címdia")
    assert title["art"]["background"] == "image", "art lives on the layout"
    agenda = next(L for L in c["layouts"] if L["name"] == "Agenda")
    assert agenda["art"]["pictures"] == 1
    assert [p["prompt"] for p in agenda["placeholders"] if p["prompt_looks_like_content"]] == [
        "01",
        "02",
        "03",
        "04",
        "05",
    ]
    num = [ch for m in c["masters"] for ch in m["chrome"] if ch["role"] == "slide_number"]
    assert num and not num[0]["placeholder"], "page number is a text box with a slidenum field"
    cards = next(P for P in c["patterns"] if P["id"] == "three-ways-we-help")
    assert cards["repeats"]["g1"]["count"] == 3 and not cards["placeholders_used"]


def test_diff_drift():
    rc, log = quiet(diff_cli.main, [str(contract_for("messy.potx")), str(TMP / "drift.potx")])
    assert rc == 1, log
    for want in (
        "'Two Content' renamed to 'Two Columns'",
        "'Picture with Caption' removed",
        "placeholder body moved",
        "accent2",
    ):
        assert want in log, (want, log)
    rc, _ = quiet(diff_cli.main, [str(contract_for("messy.potx")), str(TMP / "messy.potx")])
    assert rc == 0


def test_round_trip_18_slides():
    spec = round_trip_spec(contract_for("messy.potx"))
    rc, log = run_spec(spec, "roundtrip")
    assert rc == 0, log
    out = TMP / "roundtrip.pptx"
    rep = json.loads(out.with_suffix(".build.json").read_text())
    patterns = {r["source"] for r in rep["slides"] if r["source"].startswith("pattern:")}
    assert len(rep["slides"]) == 18 and len(patterns) >= 3, rep["slides"]
    assert not [f for f in rep["lint"] if f["severity"] == "error"], rep["lint"]
    pkg = core.Package(out)
    assert pkg.content_type(pkg.main) == core.CT_PRES, ".potx flipped to .pptx"
    assert len(pkg.slides()) == 18
    names = set(pkg.parts)
    assert (
        not [n for n in names if n.startswith("ppt/media/")]
        or len([n for n in names if n.startswith("ppt/media/")]) < 10
    ), "icon bank dropped"
    assert (
        set(pkg.reachable()) | {n for n in names if n.endswith(".rels")} | {"[Content_Types].xml"}
        >= names
    ), "no orphans"
    pres = pkg.xml(pkg.main)
    assert pres.find(core.q("p:custShowLst")) is None, "custom show of removed samples pruned"
    assert not list(pres.iter(core.q("p14:section"))), "sample sections dropped"
    for sp in pkg.slides():
        raw = pkg.parts[sp]
        if b"mc:Ignorable" in raw:
            assert b"xmlns:a14=" in raw and b"xmlns:p14=" in raw, "Ignorable prefixes stay declared"
    cards = [r["part"] for r in rep["slides"] if r["source"] == "pattern:three-ways-we-help"]
    assert b"p:timing" not in pkg.parts[cards[0]], "animation removed with the deleted card"
    texts = " ".join(
        " ".join(core.paragraphs(core.txbody(el)))
        for sp in pkg.slides()
        for _p, el, _b in core.iter_shapes(
            pkg.xml(sp).find(f"{core.q('p:cSld')}/{core.q('p:spTree')}")
        )
    )
    for leftover in ("Three ways we help", "Key numbers", "Source: fixture data", "Card heading"):
        assert leftover not in texts, leftover
    try:
        from pptx import Presentation  # dev tooling only
    except ImportError:
        print("    (python-pptx not installed: open check skipped)")
    else:
        prs = Presentation(str(out))
        assert len(prs.slides) == 18
        content = [s for i, s in enumerate(prs.slides) if i not in (0, 2, 8, 17)]
        assert all(s.has_notes_slide and s.notes_slide.notes_text_frame.text for s in content)


def test_preview_gate():
    spec = round_trip_spec(contract_for("messy.potx"))
    rc, log = run_spec(spec, "gate", "--slides", "1,10")
    assert rc == 0, log
    assert len(core.Package(TMP / "gate.preview.pptx").slides()) == 2


def test_off_contract_rejected():
    spec = {
        "contract": str(contract_for("messy.potx")),
        "slides": [
            {
                "layout": "Title Only",
                "fill": {"title": "Hi", "color": "FF00FF", "font": "Comic Sans MS"},
            },
            {"layout": "Title Only", "fill": {"title": "Hi"}, "textbox": {"x": 1, "y": 9}},
            {"layout": "Title Onyl", "fill": {"title": "Hi"}},
            {"pattern": "timeline", "fill": {"title": "T", "bogus": 1}},
        ],
    }
    rc, log = run_spec(spec, "offcontract")
    assert rc == 2 and not (TMP / "offcontract.pptx").exists(), log
    assert (
        log.count("off-contract") >= 3 and "did you mean 'Title Only'" in log and "'bogus'" in log
    ), log


def test_drift_blocks_build():
    c = json.loads(contract_for("messy.potx").read_text())
    spec = {
        "contract": str(contract_for("messy.potx")),
        "template": str(TMP / "drift.potx"),
        "slides": [{"layout": "Title Only", "fill": {"title": "x"}}],
    }
    rc, log = run_spec(spec, "drifted")
    assert rc == 3 and "BREAKING" in log and not (TMP / "drifted.pptx").exists(), log
    assert c["source"]["sha256"] != core.sha256((TMP / "drift.potx").read_bytes())


def test_overflow_and_split():
    long = [
        f"Point number {i} carries a deliberately long sentence that keeps going" for i in range(14)
    ]
    base = {
        "contract": str(contract_for("clean.pptx")),
        "slides": [
            {
                "layout": "Title and Content",
                "fill": {"title": "Too much", "body": long},
                "notes": "n",
            }
        ],
    }
    rc, log = run_spec(json.loads(json.dumps(base)), "overflow")
    assert rc == 1 and "text-overflow" in log and "needs" in log and "fits" in log, log
    base["slides"][0]["allow_split"] = True
    rc, log = run_spec(base, "split")
    assert rc == 0, log
    pkg = core.Package(TMP / "split.pptx")
    assert len(pkg.slides()) >= 2
    titles = [
        " ".join(core.paragraphs(core.txbody(el)))
        for sp in pkg.slides()
        for _p, el, _b in core.iter_shapes(
            pkg.xml(sp).find(f"{core.q('p:cSld')}/{core.q('p:spTree')}")
        )
        if core.ph_of(el) is not None and core.ph_type(core.ph_of(el)) == "title"
    ]
    assert titles[1].endswith("(cont.)"), titles


def test_edit_mode_passthrough():
    if not (TMP / "roundtrip.pptx").exists():
        test_round_trip_18_slides()
    base = TMP / "roundtrip.pptx"
    slides = [{"keep": i} for i in range(1, 19)]
    slides[3] = {"keep": 4, "fill": {"title": "Demand tripled in two regions"}}
    del slides[16]
    slides.append({"layout": "Title Only", "fill": {"title": "Added at the end"}, "notes": "new"})
    spec = {"contract": str(contract_for("messy.potx")), "base": str(base), "slides": slides}
    rc, log = run_spec(dict(spec), "edited-implicit")
    assert rc == 2 and "neither kept nor removed" in log, "implicit deletion must be refused"
    rc, log = run_spec({**spec, "output": str(base)}, "edited-overwrite")
    assert rc == 2 and "refusing to overwrite" in log, log
    spec["remove"] = [17]
    rc, log = run_spec(spec, "edited")
    assert rc == 0, log
    rep = json.loads((TMP / "edited.build.json").read_text())
    assert rep["counts"] == {"unchanged": 16, "patched": 1, "added": 1, "removed": 1}, rep["counts"]
    assert rep["unchanged_slides_byte_identical"] is True
    a, b = core.Package(base), core.Package(TMP / "edited.pptx")
    same = [p for p in a.slides() if p in b.parts and a.parts[p] == b.parts[p]]
    assert len(same) == 16


def test_replace_is_strict():
    if not (TMP / "roundtrip.pptx").exists():
        test_round_trip_18_slides()
    keep = [{"keep": i} for i in range(1, 19)]
    keep[3] = {"keep": 4, "replace": {"North grew fastest": "North grew fastest again"}}
    spec = {
        "contract": str(contract_for("messy.potx")),
        "base": str(TMP / "roundtrip.pptx"),
        "slides": keep,
    }
    rc, log = run_spec(spec, "replaced")
    assert rc == 0, log
    keep[3] = {"keep": 4, "replace": {"no such words": "x"}}
    rc, log = run_spec(spec, "replaced-missing")
    assert rc == 2 and "not found" in log, log


def test_theme_only_change_is_not_breaking():
    pkg = core.Package(TMP / "clean.pptx")
    theme = pkg.related(pkg.masters()[0], "theme")[0]
    root = pkg.xml(theme)
    acc = root.find(f".//{core.q('a:accent3')}")[0]
    acc.set("val", "123456")
    pkg.touch(theme)
    pkg.save(TMP / "clean-recolour.pptx")
    rc, log = quiet(
        diff_cli.main, [str(contract_for("clean.pptx")), str(TMP / "clean-recolour.pptx")]
    )
    assert rc == 0 and "accent3" in log and "BREAKING (0)" in log, log


def test_pattern_rename_detected():
    old = core.load_contract(contract_for("clean.pptx"))
    new = json.loads(json.dumps(old))
    p = next(x for x in new["patterns"] if x["id"] == "timeline")
    p["id"] = "roadmap"
    d = core.diff_contracts(old, new)
    assert d["breaking"] == [
        "pattern 'timeline' renamed to 'roadmap' (the sample slide's title changed; update specs)"
    ], d


def test_images_keep_aspect_and_alt():
    wide = TMP / "wide.png"
    wide.write_bytes(fixtures.png((200, 60, 60), w=40, h=10))
    spec = {
        "contract": str(contract_for("clean.pptx")),
        "slides": [
            {
                "layout": "Picture with Caption",
                "fill": {
                    "title": "Site",
                    "body": "Hall B",
                    "pic": {"image": str(wide), "alt": "Hall B floor plan"},
                },
                "notes": "n",
            }
        ],
    }
    rc, log = run_spec(spec, "picture")
    assert rc == 0, log
    pkg = core.Package(TMP / "picture.pptx")
    root = pkg.xml(pkg.slides()[0])
    rect = next(root.iter(core.q("a:srcRect")))
    assert rect.get("l") and rect.get("r"), "wide image is centre-cropped left/right, not stretched"
    assert any(c.get("descr") == "Hall B floor plan" for c in root.iter(core.q("p:cNvPr")))


def test_chart_clones_are_independent():
    c = contract_for("charts.pptx")
    chart = next(
        P
        for P in core.load_contract(c)["patterns"]
        if any(s["kind"] == "chart" for s in P["slots"].values())
    )

    def slide(vals):
        return {
            "pattern": chart["id"],
            "fill": {
                "title": f"Series {vals[0]}",
                "chart1": {
                    "categories": ["A", "B", "C"],
                    "series": [{"name": "Units", "values": vals}],
                },
            },
            "notes": "n",
        }

    rc, log = run_spec(
        {"contract": str(c), "slides": [slide([1, 2, 3]), slide([9, 8, 7])]}, "charts-out"
    )
    assert rc == 0, log
    pkg = core.Package(TMP / "charts-out.pptx")
    charts = [n for n in pkg.parts if n.startswith("ppt/charts/chart")]
    embeds = [n for n in pkg.parts if n.startswith("ppt/embeddings/")]
    assert len(charts) == 2 and len(embeds) == 2, (charts, embeds)
    vals = [[v.text for v in pkg.xml(ch).iter(core.q("c:v"))] for ch in sorted(charts)]
    assert vals[0] != vals[1] and any(v in vals[1] for v in ("9", "9.0")), vals
    for e in embeds:
        assert zipfile.ZipFile(io.BytesIO(pkg.parts[e])).testzip() is None


def test_lint_flags_hand_edits():
    if not (TMP / "roundtrip.pptx").exists():
        test_round_trip_18_slides()
    pkg = core.Package(TMP / "roundtrip.pptx")
    sp = pkg.slides()[3]
    root = pkg.xml(sp)
    run = next(root.iter(core.q("a:rPr")))
    fill = core.sub(run, "a:solidFill")
    core.sub(fill, "a:srgbClr", val="FF10F0")
    core.sub(run, "a:latin", typeface="Comic Sans MS")
    pkg.touch(sp)
    pkg.save(TMP / "hand.pptx")
    found = {
        f["code"]
        for f in lint.lint(TMP / "hand.pptx", core.load_contract(contract_for("messy.potx")))
    }
    assert {"off-palette-colour", "off-theme-font"} <= found, found


def test_render_skips_without_libreoffice():
    real = core.find_soffice
    core.find_soffice = lambda: None
    try:
        rc, log = quiet(
            render.main,
            [str(TMP / "split.pptx") if (TMP / "split.pptx").exists() else str(TMP / "clean.pptx")],
        )
        assert rc == 0 and "SKIPPED" in log, log
        rc, _ = quiet(render.main, [str(TMP / "clean.pptx"), "--require"])
        assert rc == 3
    finally:
        core.find_soffice = real


def test_powerpoint_authored_deck():
    """python-pptx's default template was saved by PowerPoint: round-trip it if available."""
    try:
        import pptx
    except ImportError:
        print("    (python-pptx not installed: skipped)")
        return
    src = Path(pptx.__file__).parent / "templates" / "default.pptx"
    out = TMP / ".pptx-style" / "pp"
    rc, log = quiet(profile_cli.main, [str(src), "--out", str(out)])
    assert rc == 0, log
    smoke = build.smoke_spec(core.load_contract(out / "contract.json"))
    smoke["contract"] = str(out / "contract.json")
    smoke["output"] = "pp-smoke.pptx"
    rc, log = run_spec(smoke, "pp-smoke")
    assert rc == 0, log
    pkg = core.Package(TMP / "pp-smoke.pptx")
    for n, raw in pkg.parts.items():
        if n.endswith(".xml") and b"mc:Ignorable=" in raw:
            root, decls = core.parse_xml(raw)
            prefixes = root.get(f"{{{core.NS['mc']}}}Ignorable", "").split()
            declared = {p for p, _u in decls}
            assert set(prefixes) <= declared, (n, prefixes, declared)


def main() -> int:
    tests = [(k, v) for k, v in globals().items() if k.startswith("test_") and callable(v)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"ok  {name}")
        except Exception as exc:  # report every failure, then exit non-zero
            failed += 1
            print(f"FAIL {name}: {type(exc).__name__}: {str(exc)[:1500]}")
    print(
        f"{'OK' if not failed else 'FAILED'}: {len(tests) - failed}/{len(tests)} passed (artifacts in {TMP})"
    )
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
